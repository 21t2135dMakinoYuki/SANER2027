from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score
import seaborn as sns
import matplotlib.pyplot as plt
import pandas as pd
import re
import json
import plotly.graph_objects as go

def RQ2(path):
    all_result = set()
    decrease_result = set()
    increase_result = set()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["result"]  
    

    for row in ws.iter_rows(min_row=2, values_only=True):
        commit_number = row[2]
        change_direction = row[6]
        UI_change_only = row[11]
        
        if UI_change_only == 0:
            continue
        
        if commit_number is None:
            continue

        if change_direction == "increased":
            increase_result.add(commit_number)
        elif change_direction == "decreased":
            decrease_result.add(commit_number)

        all_result.add(commit_number)

    unique_all_result = list({s.split('-')[0] for s in all_result})
    unique_increase_result = list({s.split('-')[0] for s in increase_result})
    unique_decrease_result = list({s.split('-')[0] for s in decrease_result})

    print("RQ2:")
    print("Decreased:" + str(len(unique_decrease_result)))
    print(unique_decrease_result)
    print("Increased:" + str(len(unique_increase_result)))
    print(unique_increase_result)
    print("Increased or Decreased" + str(len(unique_all_result)))
    print(unique_all_result)
    


def RQ3(paths, names):
  categories = [
      "Text Expansion",
      "Text Contraction",
      "Upward Shift",
      "Downward Shift",
      "Leftward Shift",
      "Rightward Shift",
      "Height Expansion",
      "Height Contraction",
      "Width Expansion",
      "Width Contraction",
      "Component Removal",
      "Component Addition",
      "Other",
  ]

  aggregated_data = {}

  for path, name in zip(paths, names):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["result"]

    for row in ws.iter_rows(values_only=True):
      test_name = row[0]
      effort = row[1]
      raw_commit = row[2]  
      change = row[6]
      UI_change_only = row[11]

      if UI_change_only == 0 or effort != "cursor_travel_distance":
        continue

      if raw_commit is not None:
        commit_number = str(raw_commit).split("-")[0].strip()
      else:
        commit_number = ""
        
      if name in [
          "glados",
          "autocannon-ui",
          "uptime-kuma",
          "monconvertisseurco2",
          "matrix",
          "uptime",
      ]:
        cc_list = [row[15], row[16], row[18], row[19]]
        tc_list = [row[12], row[13]]
      elif name == "timeoff":
        cc_list = [row[16], row[17], row[19], row[20], row[22], row[23]]
        tc_list = [row[12], row[13], row[14]]
      else:
        continue
    
      if "Other" in cc_list or "Other" in tc_list:
        continue
    
      key = (name, test_name, commit_number, change)

      if key not in aggregated_data:
        aggregated_data[key] = {
            "direction": change,
            "code_changes": set(),
            "target_changes": set(),
        }

      for cc in cc_list:
        if cc in categories:
          aggregated_data[key]["code_changes"].add(cc)
          
      for tc in tc_list:
        if tc in categories:
          aggregated_data[key]["target_changes"].add(tc)
        
  aggregated_data = {
      k: v for k, v in aggregated_data.items() 
      if v["code_changes"] and v["target_changes"]
  }
  codechanged_map = {cat: {"increased": 0, "decreased": 0} for cat in categories}
  target_map = {cat: {"increased": 0, "decreased": 0} for cat in categories}
  
  print(codechanged_map)

  for key, content in aggregated_data.items():
    dir = content["direction"]
    for cc in content["code_changes"]:
      codechanged_map[cc][dir] += 1
    for tc in content["target_changes"]:
      target_map[tc][dir] += 1

  debug_data = {
      str(k): {
          "direction": v["direction"],
          "code_changes": list(v["code_changes"]),
          "target_changes": list(v["target_changes"]),
      }
      for k, v in aggregated_data.items()
  }

  output_file = "debug_output.json"
  with open(output_file, "w", encoding="utf-8") as f:
    json.dump(debug_data, f, indent=4, ensure_ascii=False)
    
    flow_impact_target = {}
    flow_target_code = {}

    for key, content in aggregated_data.items():
        if "Other" in content["code_changes"]:
            continue

        dir_val = (
            "Increased" if content["direction"] == "increased" else "Decreased"
        )
        c_len = len(content["code_changes"])
        t_len = len(content["target_changes"])

        if c_len == 0 or t_len == 0:
            continue

        weight = 1.0 / (c_len * t_len)

        for cc in content["code_changes"]:
            cc_node = f"Code: {cc}"
            for tc in content["target_changes"]:
                tc_node = f"Target: {tc}"

                link_it = (dir_val, tc_node)
                flow_impact_target[link_it] = flow_impact_target.get(link_it, 0.0) + weight

                link_tc = (dir_val, tc_node, cc_node)
                flow_target_code[link_tc] = flow_target_code.get(link_tc, 0.0) + weight

    all_nodes_set = set()
    for (dir_val, tc_node) in flow_impact_target.keys():
        all_nodes_set.add(dir_val)
        all_nodes_set.add(tc_node)
    for (dir_val, tc_node, cc_node) in flow_target_code.keys():
        all_nodes_set.add(tc_node)
        all_nodes_set.add(cc_node)

    all_nodes = list(all_nodes_set)
    node_idx = {name: i for i, name in enumerate(all_nodes)}

    sources, targets, values, link_colors = [], [], [], []

    color_increased = "rgba(225, 29, 72, 0.45)"  
    color_decreased = "rgba(17, 202, 160, 0.45)"

    for (dir_val, tc_node), val in flow_impact_target.items():
        sources.append(node_idx[dir_val])
        targets.append(node_idx[tc_node])
        values.append(val)
        
        if dir_val == "Increased":
            link_colors.append(color_increased)
        else:
            link_colors.append(color_decreased)

    for (dir_val, tc_node, cc_node), val in flow_target_code.items():
        sources.append(node_idx[tc_node])
        targets.append(node_idx[cc_node])
        values.append(val)
        
        if dir_val == "Increased":
            link_colors.append(color_increased)
        else:
            link_colors.append(color_decreased)

    node_totals = [0] * len(all_nodes)
    for t, val in zip(targets, values):
        node_totals[t] += val
    for s, val in zip(sources, values):
        if all_nodes[s] in ["Increased", "Decreased"]:
            node_totals[s] += val

    abbr_map = {
        "Text Expansion": "TE",
        "Text Contraction": "TC",
        "Upward Shift": "US",
        "Downward Shift": "DS",
        "Leftward Shift": "LS",
        "Rightward Shift": "RS",
        "Height Expansion": "HE",
        "Height Contraction": "HC",
        "Width Expansion": "WE",
        "Width Contraction": "WC",
        "Component Removal": "CR",
        "Component Addition": "CA",
        "Other": "Other",
    }

    labeled_nodes = []
    for i, name in enumerate(all_nodes):
        clean_name = name.replace("Code: ", "").replace("Target: ", "")
        short_name = abbr_map.get(clean_name, clean_name)

        if node_totals[i] > 0:
            val = round(node_totals[i], 1)
            if val.is_integer():
                val = int(val)
            labeled_nodes.append(f"{short_name} ({val})")
        else:
            labeled_nodes.append(short_name)

    node_colors = []
    for name in all_nodes:
        if name.startswith("Code:"):
            node_colors.append("#636EFA")
        elif name.startswith("Target:"):
            node_colors.append("#EF553B")
        elif name == "Increased":
            node_colors.append("#E11D48")
        elif name == "Decreased":
            node_colors.append("#11CAA0")
        else:
            node_colors.append("#AB63FA")

    fig = go.Figure(
        data=[
            go.Sankey(
                node=dict(
                    pad=15,
                    thickness=20,
                    line=dict(color="black", width=0.5),
                    label=labeled_nodes,
                    color=node_colors,
                ),
                link=dict(
                    source=sources,
                    target=targets,
                    value=values,
                    color=link_colors,
                ),
            )
        ]
    )

    fig.update_layout(
        font_size=24,
        width=800,
        height=450,
        margin=dict(l=5, r=5, t=50, b=20),
        annotations=[
            dict(
                x=0.0,
                y=1.09,
                xref="paper",
                yref="paper",
                text="<b>Impact</b>",
                showarrow=False,
                xanchor="left",
                font=dict(size=26),
            ),
            dict(
                x=0.5,
                y=1.09,
                xref="paper",
                yref="paper",
                text="<b>Target</b>",
                showarrow=False,
                xanchor="center",
                font=dict(size=26),
            ),
            dict(
                x=1.0,
                y=1.09,
                xref="paper",
                yref="paper",
                text="<b>Code-changed</b>",
                showarrow=False,
                xanchor="right",
                font=dict(size=26),
            ),
        ],
    )
    fig.show()
    fig.write_image("sankey_diagram.pdf")
    


def kappa():
    # 0. N/A
    # 1. Text Expansion Increase in character count
    # 2. Text Contraction Decrease in character count
    # 3. Upward Shift Vertical translation toward the top of the interface
    # 4. Downward Shift Vertical translation toward the bottom of the interface
    # 5. Leftward Shift Horizontal translation toward the left of the interface
    # 6. Rightward Shift Horizontal translation toward the right of the interface
    # 7. Height Expansion Increase in vertical dimension
    # 8. Height Contraction Decrease in vertical dimension
    # 9. Width Expansion Increase in horizontal dimension
    # 10. Width Contraction Decrease in horizontal dimension
    # 11. Component Removal Deletion of an existing element from the UI structure
    # 12. Component Addition Insertion of a new element into the UI structure
    
    
    A_code = [1]*14       + [2]*5 + [5]*11 + [7]*11 + [8]*3 + [9]*20       + [10]*3 + [11]* 7 + [12]*37
    B_code = [1]*13 + [4] + [2]*5 + [5]*11 + [7]*11 + [8]*3 + [9]*19 + [8] + [10]*3 + [11]* 7 + [12]*37
    
    print("code-changed")
    print(cohen_kappa_score(A_code, B_code))
    
    A_target = [0] + [0]*3 + [3]*23 + [4]*34 + [5]*21 + [6]*20 + [9]*3 + [12]
    B_target = [9] + [4]*3 + [3]*23 + [4]*34 + [5]*21 + [6]*20 + [9]*3 + [12]

    print("target UI elements:")
    print(cohen_kappa_score(A_target, B_target))

    
def parent_child(path, name):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["result"]  
    
    unique_pairs = set()
    unique_indirect_pairs = set()
    
    for row in ws.iter_rows(values_only=True):
        if row[0] is not None and row[1] == "cursor_travel_distance" and row[11]==1:
            test_name = row[0] 
            commit_number = row[2]
            
            unique_pairs.add((test_name, commit_number))
            
            if name == "timeoff":
                if row[24] == 1:
                    unique_indirect_pairs.add((test_name, commit_number))
            elif name == "glados":
                if row[20] == 1:
                    unique_indirect_pairs.add((test_name, commit_number))
    print(name)
    print("total cases:")
    print(len(unique_pairs))
    if name=="timeoff":
        print("Timeoff: 4 'Other' cases identified")
    print("indirect change cases:")
    print(len(unique_indirect_pairs))
    
